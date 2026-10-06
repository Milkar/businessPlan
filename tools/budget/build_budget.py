import sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter as L
from openpyxl.comments import Comment

OUT = sys.argv[1]
FONT = "Arial"
BLUE = Font(name=FONT, color="0000FF")
BLACK = Font(name=FONT, color="000000")
GREEN = Font(name=FONT, color="008000")
BOLD = Font(name=FONT, bold=True)
TITLE = Font(name=FONT, bold=True, size=14)
HDR = Font(name=FONT, bold=True, color="FFFFFF")
HDRFILL = PatternFill("solid", fgColor="1F3864")
SECFILL = PatternFill("solid", fgColor="D9E1F2")
YELLOW = PatternFill("solid", fgColor="FFFF00")
TOTFILL = PatternFill("solid", fgColor="F2F2F2")
NUM = '#,##0.0;(#,##0.0);"-"'
INT = '#,##0;(#,##0);"-"'
PCT = '0.0%;(0.0%);"-"'
thin = Side(style="thin", color="BFBFBF")
TOP = Border(top=Side(style="thin", color="000000"))

wb = Workbook()

def style_all(ws):
    for row in ws.iter_rows():
        for c in row:
            if c.font is None or c.font.name != FONT:
                f = c.font
                c.font = Font(name=FONT, bold=f.bold, color=f.color, size=f.size, italic=f.italic)

def header(ws, r, labels, start=1):
    for i, t in enumerate(labels):
        c = ws.cell(r, start + i, t)
        c.font = HDR; c.fill = HDRFILL
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

# ---------------------------------------------------------------- 说明
ws0 = wb.active; ws0.title = "说明"
lines = [
    ("AIVC / 知微 Chronoid — 5000 万元融资 12 个月预算模拟", TITLE),
    ("单位：人民币万元。本模型为测算工具，全部假设需团队与财务顾问核对后才能对外引用。", None),
    ("", None),
    ("如何使用", BOLD),
    ("1. 只改「假设」表和「人员」表里的蓝色数字（黄色底 = 关键假设）。其他表全部为公式，会自动重算。", None),
    ("2. 「月度预算」：Plan A 逐月支出，按六大类汇总。", None),
    ("3. 「汇总」：资金分配、准备金、单例成本、里程碑节点的累计支出。", None),
    ("4. 「PlanB情景」：准备金动用规则、三个情景下的现金余额与跑道月数。", None),
    ("", None),
    ("颜色约定", BOLD),
    ("蓝色字 = 手动输入的假设；黑色字 = 公式；绿色字 = 引用其他工作表；黄色底 = 关键假设，请优先核对。", None),
    ("", None),
    ("口径说明", BOLD),
    ("· 12 个月内不计任何收入（保守口径）：首个付费项目即使落地，也只作为上行空间，不抵减支出。", None),
    ("· 病例成本在入组当月一次性计入（简化）；实际上患者 T1（复发）数据可能晚于 12 个月产生。", None),
    ("· 不含增值税、汇率波动、股权激励费用；设备按采购当期现金支出计（不做折旧）。", None),
    ("· 价格基于国内一线城市生物医药初创企业的常见市场价区间估算，未经报价核实。【待核】", None),
]
for i, (t, f) in enumerate(lines, 1):
    c = ws0.cell(i, 1, t)
    if f: c.font = f
ws0.column_dimensions["A"].width = 110

# ---------------------------------------------------------------- 假设
wa = wb.create_sheet("假设")
wa["A1"] = "关键假设（蓝色可改）"; wa["A1"].font = TITLE
header(wa, 3, ["参数", "数值", "单位", "说明 / 依据"])
A = {}  # name -> absolute ref
r = 4
def sec(title):
    global r
    for col in range(1, 5):
        wa.cell(r, col).fill = SECFILL
    wa.cell(r, 1, title).font = BOLD
    r += 1
def inp(name, label, val, unit, note, fmt=NUM, key=False):
    global r
    wa.cell(r, 1, label)
    c = wa.cell(r, 2, val); c.font = BLUE; c.number_format = fmt
    if key: c.fill = YELLOW
    wa.cell(r, 3, unit); wa.cell(r, 4, note)
    A[name] = f"假设!$B${r}"
    r += 1

sec("融资与准备金")
inp("raise", "本轮融资额", 5000, "万元", "用户口径：5000 万元覆盖 12 个月", NUM, True)
inp("res_target", "Plan B 准备金目标比例", 0.20, "%", "建议 20%–25%：生物医药早期项目的融资周期通常 6–9 个月，需留出过桥资金", PCT, True)

sec("实验室与设备")
inp("capex", "设备采购总额", 800, "万元", "高内涵成像、长时程活细胞成像、培养箱 / 生物安全柜 / 离心机等基础设备【待核：以报价为准】", NUM, True)
inp("capex_m1", "设备采购：第 1 个月占比", 0.5, "%", "", PCT)
inp("capex_m2", "设备采购：第 2 个月占比", 0.3, "%", "", PCT)
inp("capex_m3", "设备采购：第 3 个月占比", 0.2, "%", "三个月占比之和应为 100%", PCT)
inp("fitout", "实验室装修（第 1 个月一次性）", 150, "万元", "约 500 m² P2 级细胞实验室", NUM)
inp("lab_area", "实验室面积", 500, "m²", "", INT)
inp("lab_rent", "实验室租金", 4.5, "元/m²/天", "园区实验室租金区间估算【待核】", NUM)
inp("consum", "通用实验耗材（第 2 个月起）", 8, "万元/月", "与具体病例无关的耗材、试剂", NUM)
inp("maint", "设备维保费率（第 4 个月起）", 0.05, "%/年", "按设备总额计", PCT)

sec("配对病例单例成本")
inp("success", "类器官建模成功率", 0.70, "%", "失败病例只发生「样本获取 + 建模」成本【待补充：团队历史成功率】", PCT, True)
inp("c_acq", "样本获取与临床随访", 1.5, "万元/例", "采样、病理、医院端数据整理与随访", NUM)
inp("c_est", "类器官建立与扩增", 2.0, "万元/例", "培养基、基质胶、建库与质控", NUM)
inp("c_pert", "扰动实验（多药物 × 多剂量 × 多时间点）", 3.0, "万元/成功例", "统一 SOP 下的扰动面板", NUM)
inp("c_sc", "单细胞测序（4 个样本/例）", 4.0, "万元/成功例", "患者 T0 / PDO T0 / PDO T1 / 患者 T1，约 1 万元/样本", NUM)
inp("c_bulk", "bulk RNA + WES", 1.7, "万元/成功例", "bulk RNA 约 8 条件 × 0.1；WES 约 3 样本 × 0.3", NUM)
inp("c_omics", "蛋白组 + 代谢组", 1.6, "万元/成功例", "约 4 样本 × 0.4", NUM)
inp("c_spatial", "空间转录组（按 50% 病例做）", 1.5, "万元/成功例", "约 3 万元/样本 × 50%", NUM)
inp("c_img", "成像与功能表型", 0.8, "万元/成功例", "高内涵 / 活细胞成像运行耗材", NUM)
inp("retro_n", "回顾性验证样本数（第 8–11 个月）", 30, "例", "利用已有样本库做药物响应回顾性验证（第 12 个月里程碑）", INT)
inp("retro_c", "回顾性验证单例成本", 3.0, "万元/例", "", NUM)

sec("数据与算力")
inp("gpu_start", "云 GPU 开始月份", 2, "月", "", INT)
inp("gpu_full", "云 GPU 满负荷月份", 5, "月", "之前按 50% 计", INT)
inp("gpu", "云 GPU 满负荷月费", 15, "万元/月", "约 8 卡训练集群的云租用【待核】", NUM, True)
inp("dp_setup", "数据平台搭建（第 2 个月一次性）", 60, "万元", "LIMS、数据湖，扩展现有 Verqura 系统", NUM)
inp("dp_run", "数据平台运维（第 3 个月起）", 3, "万元/月", "存储、备份、等保", NUM)
inp("sw", "软件许可", 40, "万元/年", "分析软件、图像处理、协作工具", NUM)
inp("ext_data", "外部数据 / 数据库授权（第 4 个月一次性）", 30, "万元", "", NUM)

sec("临床合作、合规与知识产权")
inp("centers", "合作医院数", 4, "家", "【待补充：合作中心名单与协议状态】", INT, True)
inp("center_fee", "单中心项目年费", 40, "万元/年", "含伦理审查、数据管理、研究护士支持", NUM)
inp("hgr", "人类遗传资源审批 / 伦理合规咨询（第 1–3 个月）", 40, "万元", "遗传办审批或备案、知情同意文件", NUM)
inp("datasec", "数据安全与出境合规（第 4 个月）", 30, "万元", "数据出境安全评估 / 标准合同、等保", NUM)
inp("patent_n", "专利申请数", 8, "件", "国内发明 + PCT", INT)
inp("patent_c", "单件专利费用", 6, "万元/件", "", NUM)
inp("legal", "法律、审计、财务服务", 40, "万元/年", "", NUM)

sec("商务与运营")
inp("office_area", "办公面积", 300, "m²", "", INT)
inp("office_rent", "办公租金", 4.0, "元/m²/天", "", NUM)
inp("travel", "差旅（含中德往返）", 4, "万元/月", "", NUM)
inp("mkt", "学术会议与市场", 40, "万元/年", "", NUM)
inp("pilot", "BD 试点项目投入（第 7–12 个月）", 100, "万元", "为争取首个付费项目做的验证性试点", NUM)
inp("admin", "行政、IT、保险杂费", 2, "万元/月", "", NUM)
inp("contg", "Plan A 内部不可预见费率", 0.03, "%", "按当月其他支出计；与 Plan B 准备金分开", PCT)
inp("recruit", "招聘费率", 0.10, "%", "按新增人员年度人力成本计（猎头 / 招聘渠道）", PCT)

wa.column_dimensions["A"].width = 44; wa.column_dimensions["B"].width = 12
wa.column_dimensions["C"].width = 14; wa.column_dimensions["D"].width = 80
wa.freeze_panes = "A4"

# ---------------------------------------------------------------- 人员
wp = wb.create_sheet("人员")
wp["A1"] = "人员计划（月末在岗人数；蓝色可改）"; wp["A1"].font = TITLE
wp["A2"] = "年度人力成本 = 年薪 + 社保公积金等（单位：万元/人/年）。M0 列固定为 0，用于计算第 1 个月的新增人数。"
header(wp, 4, ["岗位", "年度人力成本\n(万元/人)", "M0"] + [f"M{m}" for m in range(1, 13)] + ["人月合计", "12 个月成本\n(万元)"])
roles = [
    ("创始团队（CEO / CTO / ML 负责人 / 运营）", 45, [4]*12, "创始人取市场价下限"),
    ("资深科学家（类器官 / 肿瘤生物学）", 55, [1, 1, 2, 2, 2, 2, 2, 2, 2, 2, 2, 2], ""),
    ("实验技术员 / 研究助理", 18, [2, 4, 4, 6, 6, 6, 8, 8, 10, 10, 10, 10], ""),
    ("生物信息 / 数据工程师", 40, [0, 1, 2, 2, 3, 3, 3, 3, 3, 3, 3, 3], ""),
    ("机器学习工程师", 55, [0, 1, 1, 2, 2, 3, 3, 3, 4, 4, 4, 4], ""),
    ("临床研究协调员（CRC）", 15, [0, 1, 2, 2, 3, 3, 3, 3, 3, 3, 3, 3], ""),
    ("质量 / 注册与合规", 40, [0, 0, 0, 1, 1, 1, 1, 1, 1, 1, 1, 1], ""),
    ("商务拓展（BD）", 45, [0, 0, 0, 1, 1, 1, 2, 2, 2, 2, 2, 2], ""),
    ("行政 / 财务 / HR", 20, [1, 1, 2, 2, 2, 2, 2, 2, 2, 2, 2, 2], ""),
]
r0 = 5
for i, (name, cost, hc, note) in enumerate(roles):
    rr = r0 + i
    wp.cell(rr, 1, name)
    c = wp.cell(rr, 2, cost); c.font = BLUE; c.number_format = NUM
    wp.cell(rr, 3, 0).font = BLUE
    for m in range(12):
        c = wp.cell(rr, 4 + m, hc[m]); c.font = BLUE; c.number_format = INT
    wp.cell(rr, 16, f"=SUM(D{rr}:O{rr})").number_format = INT
    wp.cell(rr, 17, f"=B{rr}*P{rr}/12").number_format = NUM
r1 = r0 + len(roles) - 1
rt = r1 + 1
wp.cell(rt, 1, "在岗人数合计").font = BOLD
for col in range(3, 17):
    c = wp.cell(rt, col, f"=SUM({L(col)}{r0}:{L(col)}{r1})"); c.number_format = INT; c.font = BOLD
wp.cell(rt, 17, f"=SUM(Q{r0}:Q{r1})").number_format = NUM
rpay = rt + 1
wp.cell(rpay, 1, "当月人力成本（万元）").font = BOLD
for col in range(4, 16):
    c = wp.cell(rpay, col, f"=SUMPRODUCT($B${r0}:$B${r1},{L(col)}{r0}:{L(col)}{r1})/12"); c.number_format = NUM
wp.cell(rpay, 17, f"=SUM(D{rpay}:O{rpay})").number_format = NUM
rrec = rpay + 1
wp.cell(rrec, 1, "当月招聘费用（万元）").font = BOLD
for col in range(4, 16):
    cur, prev = L(col), L(col - 1)
    c = wp.cell(rrec, col,
        f"={A['recruit']}*SUMPRODUCT($B${r0}:$B${r1},({cur}{r0}:{cur}{r1}>{prev}{r0}:{prev}{r1})*({cur}{r0}:{cur}{r1}-{prev}{r0}:{prev}{r1}))")
    c.number_format = NUM
wp.cell(rrec, 17, f"=SUM(D{rrec}:O{rrec})").number_format = NUM
wp.column_dimensions["A"].width = 40; wp.column_dimensions["B"].width = 14
for col in range(3, 18):
    wp.column_dimensions[L(col)].width = 8
wp.column_dimensions["P"].width = 10; wp.column_dimensions["Q"].width = 12
wp.freeze_panes = "C5"
PAY = rpay; REC = rrec; HC = rt

# ---------------------------------------------------------------- 月度预算
wm = wb.create_sheet("月度预算")
wm["A1"] = "Plan A 月度预算（万元）"; wm["A1"].font = TITLE
wm["A2"] = "第 4 行为月份序号；第 5 行病例入组数为输入（蓝色）。其余为公式。"
header(wm, 3, ["类别", "项目"] + [f"M{m}" for m in range(1, 13)] + ["12 个月合计"])
wm.cell(4, 2, "月份序号")
for m in range(12):
    c = wm.cell(4, 3 + m, m + 1); c.number_format = INT
wm.cell(5, 2, "新入组病例数（例）")
cases = [0, 0, 2, 3, 4, 5, 8, 8, 9, 9, 9, 9]
for m in range(12):
    c = wm.cell(5, 3 + m, cases[m]); c.font = BLUE; c.fill = YELLOW; c.number_format = INT
wm.cell(5, 15, "=SUM(C5:N5)").number_format = INT
wm.cell(6, 2, "其中成功建模（例）")
for m in range(12):
    c = wm.cell(6, 3 + m, f"={L(3+m)}5*{A['success']}"); c.number_format = NUM
wm.cell(6, 15, "=SUM(C6:N6)").number_format = NUM

ROW = {}
rr = 8
def mon(f):  # f(colLetter, monthCell) -> formula
    return [f(L(3 + m), f"{L(3+m)}$4") for m in range(12)]
CATS = []
def category(cat, items):
    global rr
    start = rr
    for key, label, fn in items:
        wm.cell(rr, 1, cat if rr == start else "")
        wm.cell(rr, 2, label)
        for m, fx in enumerate(mon(fn)):
            c = wm.cell(rr, 3 + m, fx); c.number_format = NUM
            if "人员!" in fx and fx.count("!") == 1 and fx.startswith("=人员!"):
                c.font = GREEN
        wm.cell(rr, 15, f"=SUM(C{rr}:N{rr})").number_format = NUM
        ROW[key] = rr
        rr += 1
    end = rr - 1
    wm.cell(rr, 2, f"{cat} 小计").font = BOLD
    for col in range(3, 16):
        c = wm.cell(rr, col, f"=SUM({L(col)}{start}:{L(col)}{end})"); c.number_format = NUM; c.font = BOLD
        c.fill = TOTFILL
    wm.cell(rr, 1).fill = TOTFILL; wm.cell(rr, 2).fill = TOTFILL
    ROW["sub_" + cat] = rr
    CATS.append(cat)
    rr += 2

pcol = lambda col: L(ord(col) - ord("C") + 4)  # 月度预算 C->人员 D

category("人员", [
    ("pay", "人员薪酬（含社保公积金）", lambda c, m: f"=人员!{pcol(c)}{PAY}"),
    ("rec", "招聘费用", lambda c, m: f"=人员!{pcol(c)}{REC}"),
])
category("实验与样本", [
    ("capex", "设备采购", lambda c, m: f"={A['capex']}*(({m}=1)*{A['capex_m1']}+({m}=2)*{A['capex_m2']}+({m}=3)*{A['capex_m3']})"),
    ("fit", "实验室装修", lambda c, m: f"=({m}=1)*{A['fitout']}"),
    ("labrent", "实验室租金", lambda c, m: f"={A['lab_area']}*{A['lab_rent']}*30/10000"),
    ("caseA", "样本获取与类器官建立（全部入组病例）", lambda c, m: f"={c}$5*({A['c_acq']}+{A['c_est']})"),
    ("caseB", "扰动实验与多组学读出（成功建模病例）", lambda c, m: f"={c}$6*({A['c_pert']}+{A['c_sc']}+{A['c_bulk']}+{A['c_omics']}+{A['c_spatial']}+{A['c_img']})"),
    ("retro", "回顾性验证样本", lambda c, m: f"=AND({m}>=8,{m}<=11)*{A['retro_n']}*{A['retro_c']}/4"),
    ("cons", "通用实验耗材", lambda c, m: f"=({m}>=2)*{A['consum']}"),
    ("mnt", "设备维保", lambda c, m: f"=({m}>=4)*{A['capex']}*{A['maint']}/12"),
])
category("数据与算力", [
    ("gpu", "云 GPU 算力", lambda c, m: f"=IF({m}<{A['gpu_start']},0,IF({m}<{A['gpu_full']},0.5,1))*{A['gpu']}"),
    ("dps", "数据平台搭建", lambda c, m: f"=({m}=2)*{A['dp_setup']}"),
    ("dpr", "数据平台运维", lambda c, m: f"=({m}>=3)*{A['dp_run']}"),
    ("sw", "软件许可", lambda c, m: f"={A['sw']}/12"),
    ("ext", "外部数据 / 数据库授权", lambda c, m: f"=({m}=4)*{A['ext_data']}"),
])
category("临床合作、合规与IP", [
    ("ctr", "合作医院项目费", lambda c, m: f"=({m}>=2)*{A['centers']}*{A['center_fee']}/12"),
    ("hgr", "人类遗传资源 / 伦理合规", lambda c, m: f"=({m}<=3)*{A['hgr']}/3"),
    ("dsec", "数据安全与出境合规", lambda c, m: f"=({m}=4)*{A['datasec']}"),
    ("pat", "专利申请", lambda c, m: f"={A['patent_n']}*{A['patent_c']}/12"),
    ("leg", "法律、审计、财务", lambda c, m: f"={A['legal']}/12"),
])
category("商务与运营", [
    ("off", "办公租金", lambda c, m: f"={A['office_area']}*{A['office_rent']}*30/10000"),
    ("trv", "差旅", lambda c, m: f"={A['travel']}"),
    ("mkt", "学术会议与市场", lambda c, m: f"={A['mkt']}/12"),
    ("pil", "BD 试点项目投入", lambda c, m: f"=({m}>=7)*{A['pilot']}/6"),
    ("adm", "行政、IT、保险杂费", lambda c, m: f"={A['admin']}"),
])
# contingency: rate * sum of the other subtotals
subs = [ROW["sub_" + c] for c in CATS]
wm.cell(rr, 1, "不可预见费")
wm.cell(rr, 2, "Plan A 内部不可预见费")
for m in range(12):
    col = L(3 + m)
    c = wm.cell(rr, 3 + m, f"={A['contg']}*(" + "+".join(f"{col}{s}" for s in subs) + ")"); c.number_format = NUM
wm.cell(rr, 15, f"=SUM(C{rr}:N{rr})").number_format = NUM
ROW["contg"] = rr; ROW["sub_不可预见费"] = rr
rr += 2
TOTAL = rr
wm.cell(rr, 1, "合计").font = BOLD
wm.cell(rr, 2, "Plan A 当月支出").font = BOLD
for col in range(3, 16):
    c = wm.cell(rr, col, f"=" + "+".join(f"{L(col)}{s}" for s in subs + [ROW['contg']]))
    c.number_format = NUM; c.font = BOLD; c.border = TOP
rr += 1
CUM = rr
wm.cell(rr, 2, "Plan A 累计支出").font = BOLD
for m in range(12):
    col = L(3 + m)
    c = wm.cell(rr, 3 + m, f"=SUM($C${TOTAL}:{col}{TOTAL})"); c.number_format = NUM
rr += 1
CASHA = rr
wm.cell(rr, 2, "账面现金余额（含准备金）").font = BOLD
for m in range(12):
    col = L(3 + m)
    c = wm.cell(rr, 3 + m, f"={A['raise']}-{col}{CUM}"); c.number_format = NUM
wm.column_dimensions["A"].width = 18; wm.column_dimensions["B"].width = 36
for col in range(3, 16):
    wm.column_dimensions[L(col)].width = 9
wm.column_dimensions["O"].width = 12
wm.freeze_panes = "C4"

# ---------------------------------------------------------------- 汇总
wsum = wb.create_sheet("汇总")
wsum["A1"] = "预算汇总与准备金（万元）"; wsum["A1"].font = TITLE
header(wsum, 3, ["项目", "金额（万元）", "占融资额", "占 Plan A 支出", "对应里程碑 / 用途"])
mil = {
    "人员": "团队按飞轮环节逐步扩充（人数见「人员」表）",
    "实验与样本": "第 6 个月首批配对队列；第 12 个月回顾性验证 + 前瞻性评分启动",
    "数据与算力": "AIVC 模型训练与数据平台；第 12 个月具备 v1 模型训练条件",
    "临床合作、合规与IP": "合作中心、遗传资源与伦理合规、数据出境合规、专利布局",
    "商务与运营": "首个付费 / 联合开发项目的试点与商务拓展",
    "不可预见费": "Plan A 内部的小额超支缓冲（与准备金分开）",
}
r = 4
first = r
for cat in CATS + ["不可预见费"]:
    wsum.cell(r, 1, cat)
    c = wsum.cell(r, 2, f"=月度预算!O{ROW['sub_' + cat]}"); c.font = GREEN; c.number_format = NUM
    wsum.cell(r, 3, f"=B{r}/{A['raise']}").number_format = PCT
    wsum.cell(r, 4, f"=B{r}/$B${first + len(CATS) + 1}").number_format = PCT
    wsum.cell(r, 5, mil[cat])
    r += 1
PA = r
wsum.cell(r, 1, "Plan A 12 个月支出合计").font = BOLD
c = wsum.cell(r, 2, f"=SUM(B{first}:B{r-1})"); c.number_format = NUM; c.font = BOLD; c.border = TOP
wsum.cell(r, 3, f"=B{r}/{A['raise']}").number_format = PCT
wsum.cell(r, 4, f"=B{r}/B{r}").number_format = PCT
r += 1
RES = r
wsum.cell(r, 1, "Plan B 准备金（独立账户）").font = BOLD
c = wsum.cell(r, 2, f"={A['raise']}-B{PA}"); c.number_format = NUM; c.font = BOLD; c.fill = YELLOW
wsum.cell(r, 3, f"=B{r}/{A['raise']}").number_format = PCT
wsum.cell(r, 5, "仅在触发条件出现、经董事会批准后动用（见 PlanB情景）")
r += 1
wsum.cell(r, 1, "本轮融资额").font = BOLD
c = wsum.cell(r, 2, f"={A['raise']}"); c.number_format = NUM; c.font = GREEN
wsum.cell(r, 3, f"=B{r}/{A['raise']}").number_format = PCT
r += 2
wsum.cell(r, 1, "检查：准备金是否达到目标比例").font = BOLD
wsum.cell(r, 2, f"=IF(C{RES}>={A['res_target']},\"达标\",\"未达标：需压缩 Plan A\")")
wsum.cell(r, 3, f"={A['res_target']}").number_format = PCT
wsum.cell(r, 4, "目标比例")
CHECK = r
r += 2

wsum.cell(r, 1, "关键指标").font = BOLD; r += 1
K = {}
def kpi(key, label, fx, fmt=NUM, note=""):
    global r
    wsum.cell(r, 1, label)
    c = wsum.cell(r, 2, fx); c.number_format = fmt
    wsum.cell(r, 5, note)
    K[key] = r
    r += 1
kpi("cases", "12 个月入组病例", f"=月度预算!O5", INT)
kpi("succ", "其中成功建模（配对病例）", f"=月度预算!O6", NUM)
kpi("hc12", "第 12 个月在岗人数", f"=人员!O{HC}", INT)
kpi("cpc_dir", "单例直接实验成本（按成功病例分摊）", f"=(月度预算!O{ROW['caseA']}+月度预算!O{ROW['caseB']})/B{r-2}", NUM, "样本获取 + 建模 + 扰动 + 多组学，含失败病例的分摊")
kpi("burn_avg", "Plan A 月均支出", f"=B{PA}/12", NUM)
kpi("burn_q4", "Plan A 第 10–12 个月月均支出", f"=AVERAGE(月度预算!L{TOTAL}:N{TOTAL})", NUM, "进入稳态后的月度消耗，用于估算下一轮融资规模")
kpi("cum6", "第 6 个月累计支出（Level 1 闸门）", f"=月度预算!H{CUM}", NUM, "闸门未过时，剩余资金 = 融资额 − 本行")
kpi("cum12", "第 12 个月累计支出", f"=月度预算!N{CUM}", NUM)
kpi("next", "下一轮融资参考规模（稳态月耗 × 24 个月）", f"=B{K['burn_q4']}*24", NUM, "仅为量级参考：下一轮通常覆盖 18–24 个月且团队会扩大")
wsum.column_dimensions["A"].width = 40; wsum.column_dimensions["B"].width = 14
wsum.column_dimensions["C"].width = 11; wsum.column_dimensions["D"].width = 14
wsum.column_dimensions["E"].width = 70

# ---------------------------------------------------------------- PlanB情景
wb_ = wb.create_sheet("PlanB情景")
w = wb_
w["A1"] = "Plan B：准备金动用规则与情景测算（万元）"; w["A1"].font = TITLE
r = 3
w.cell(r, 1, "一、维持期月度支出（Plan B 最低运转成本，蓝色可改）").font = BOLD; r += 1
header(w, r, ["项目", "数值", "单位", "说明"]); r += 1
B = {}
def binp(key, label, val, unit, note, fmt=NUM):
    global r
    w.cell(r, 1, label)
    c = w.cell(r, 2, val); c.font = BLUE; c.number_format = fmt
    w.cell(r, 3, unit); w.cell(r, 4, note)
    B[key] = f"$B${r}"; r += 1
binp("keep_n", "保留人员数", 14, "人", "创始团队 + 关键科学家 + 核心实验 / 数据 / 模型人员", INT)
binp("keep_c", "保留人员人均年度人力成本", 40, "万元/人/年", "")
binp("keep_lab", "最低实验支出（耗材 + 维持建模）", 15, "万元/月", "保持类器官库存活与少量病例入组")
binp("keep_gpu", "最低算力与数据平台", 5, "万元/月", "")
binp("keep_ops", "最低运营（租金之外的行政、差旅、合规）", 8, "万元/月", "")
w.cell(r, 1, "实验室 + 办公租金")
c = w.cell(r, 2, f"=({A['lab_area']}*{A['lab_rent']}+{A['office_area']}*{A['office_rent']})*30/10000"); c.number_format = NUM; c.font = GREEN
w.cell(r, 3, "万元/月"); B["rent"] = f"$B${r}"; r += 1
w.cell(r, 1, "维持期月度支出合计").font = BOLD
c = w.cell(r, 2, f"={B['keep_n']}*{B['keep_c']}/12+{B['keep_lab']}+{B['keep_gpu']}+{B['keep_ops']}+{B['rent']}")
c.number_format = NUM; c.font = BOLD; c.fill = YELLOW
w.cell(r, 3, "万元/月"); MB = f"$B${r}"; r += 2

w.cell(r, 1, "二、Plan B-1 削减参数（科学闸门未过时启用，蓝色可改）").font = BOLD; r += 1
header(w, r, ["项目", "数值", "单位", "说明"]); r += 1
binp("trig", "触发月份（从该月起执行削减）", 7, "月", "第 6 个月 Level 1 闸门评估未通过", INT)
binp("cut_lab", "实验与样本支出削减比例", 0.5, "%", "病例入组减半，集中在建模成功率最高的瘤种", PCT)
binp("cut_data", "数据与算力支出削减比例", 0.3, "%", "", PCT)
binp("cut_ops", "商务与运营支出削减比例", 0.5, "%", "暂停 BD 试点，压缩会议与差旅", PCT)
w.cell(r, 1, "人员：冻结招聘，维持触发前一个月的人力成本"); r += 2

w.cell(r, 1, "三、逐月现金测算").font = BOLD; r += 1
header(w, r, ["情景", "项目"] + [f"M{m}" for m in range(1, 13)] + ["12 个月合计"]); r += 1
mrow = r
w.cell(r, 2, "月份序号")
for m in range(12):
    w.cell(r, 3 + m, m + 1).number_format = INT
r += 1
def mb(key):  # 月度预算 row ref for column letter
    return ROW[key]
# Scenario A
SA = r
w.cell(r, 1, "情景 A：Plan A 正常推进"); w.cell(r, 2, "当月支出")
for m in range(12):
    c = w.cell(r, 3 + m, f"=月度预算!{L(3+m)}{TOTAL}"); c.font = GREEN; c.number_format = NUM
w.cell(r, 15, f"=SUM(C{r}:N{r})").number_format = NUM
r += 1
SB = r
w.cell(r, 1, "情景 B：第 6 个月闸门未过，第 7 个月起削减"); w.cell(r, 2, "当月支出")
sub = lambda cat, col: f"月度预算!{col}{ROW['sub_' + cat]}"
for m in range(12):
    col = L(3 + m)
    mc = f"{col}${mrow}"
    pay_fixed = f"INDEX(月度预算!$C${ROW['pay']}:$N${ROW['pay']},{B['trig']}-1)"
    planb = (f"{pay_fixed}+{sub('实验与样本', col)}*(1-{B['cut_lab']})+{sub('数据与算力', col)}*(1-{B['cut_data']})"
             f"+{sub('临床合作、合规与IP', col)}+{sub('商务与运营', col)}*(1-{B['cut_ops']})")
    c = w.cell(r, 3 + m, f"=IF({mc}<{B['trig']},月度预算!{col}{TOTAL},{planb})"); c.number_format = NUM
w.cell(r, 15, f"=SUM(C{r}:N{r})").number_format = NUM
r += 1
SC = r
w.cell(r, 1, "情景 C：闸门未过 + 下一轮融资受阻，第 7 个月起直接进入维持期"); w.cell(r, 2, "当月支出")
for m in range(12):
    col = L(3 + m)
    c = w.cell(r, 3 + m, f"=IF({col}${mrow}<{B['trig']},月度预算!{col}{TOTAL},{MB})"); c.number_format = NUM
w.cell(r, 15, f"=SUM(C{r}:N{r})").number_format = NUM
r += 1
for s, lab in ((SA, "A"), (SB, "B"), (SC, "C")):
    w.cell(r, 1, f"情景 {lab}"); w.cell(r, 2, "账面现金余额")
    for m in range(12):
        col = L(3 + m)
        c = w.cell(r, 3 + m, f"={A['raise']}-SUM($C${s}:{col}{s})"); c.number_format = NUM
    r += 1
r += 1

w.cell(r, 1, "四、情景结果").font = BOLD; r += 1
header(w, r, ["情景", "12 个月支出", "第 12 个月末现金", "之后按维持期可再支撑（月）", "总跑道（月）", "说明"]); r += 1
notes = {
    "A": "准备金未动用，作为下一轮融资的过桥资金",
    "B": "科学风险出现时收缩，仍完成核心验证，跑道延长",
    "C": "最坏情况：保住团队、类器官库与数据资产，等待融资或战略合作",
}
for s, lab, name in ((SA, "A", "情景 A：Plan A 正常推进"), (SB, "B", "情景 B：闸门未过，收缩"), (SC, "C", "情景 C：闸门未过 + 融资受阻")):
    w.cell(r, 1, name)
    w.cell(r, 2, f"=O{s}").number_format = NUM
    w.cell(r, 3, f"={A['raise']}-O{s}").number_format = NUM
    w.cell(r, 4, f"=IF({MB}>0,C{r}/{MB},0)").number_format = NUM
    w.cell(r, 5, f"=12+D{r}").number_format = NUM
    w.cell(r, 6, notes[lab])
    r += 1
r += 1
w.cell(r, 1, "五、准备金动用规则（建议写入股东协议 / 董事会议事规则）").font = BOLD; r += 1
rules = [
    ("触发条件 1", "第 6 个月 Level 1 闸门（重复性、周转时间、失败率、培养偏差模型）未达阈值 → 执行情景 B"),
    ("触发条件 2", "第 9 个月仍未拿到下一轮投资意向书（TS），且现金跑道 < 6 个月 → 冻结招聘，准备进入维持期"),
    ("触发条件 3", "第 12 个月首个付费 / 联合开发项目未落地 → 转向类器官检测与研究服务，以服务收入延长跑道"),
    ("动用程序", "准备金存放于独立账户；动用需董事会（含投资人董事）批准，并同步公布收缩方案"),
    ("不得动用于", "Plan A 的常规超支（由 Plan A 内部不可预见费覆盖）、新增非核心岗位、未经批准的设备采购"),
]
for k, v in rules:
    w.cell(r, 1, k).font = BOLD; w.cell(r, 2, v); r += 1
w.column_dimensions["A"].width = 46; w.column_dimensions["B"].width = 16
for col in range(3, 15):
    w.column_dimensions[L(col)].width = 10
w.column_dimensions["O"].width = 12
w.column_dimensions["D"].width = 14; w.column_dimensions["F"].width = 12

for s in wb.worksheets:
    style_all(s)
wb.save(OUT)
print("saved", OUT)

# fullCalcOnLoad so Excel / WPS always recompute on open
from openpyxl import load_workbook as _lw
_wb = _lw(OUT); _wb.calculation.fullCalcOnLoad = True; _wb.save(OUT)
